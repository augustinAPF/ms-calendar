"""
Field MeritTrac Dashboard (desk page /app/field-merittrac-dashboard).

Brings "Field Meritrac Test URL" (test links sent) and "Field MeritTrac Test
Result" (results MeritTrac posted back) together, one row per applicant:
link sent → test window → attempted → score band → evaluator review →
Registration Form status.

Every endpoint here is permission-scoped on the server:
- links, results and applicants are read with frappe.get_list (applies the
  user's DocType and User Permissions), never frappe.get_all — a doctype the
  user can't read simply contributes nothing;
- child rows (section scores) and Field Evaluator Marks are only looked up
  for the results get_list already returned for this user;
- the candidate detail endpoint checks read permission on that exact
  Field Registration Form record, and only returns links / results the user
  can read.
The page computes cards, tables, drilldowns and CSV exports from the rows
get_dashboard_data returns, and the Excel download (download_data) runs the
same _build, so an export can never be wider than what the user is allowed
to see on screen. "Full data" respects field-level permissions too.
"""

from datetime import datetime

import frappe
from frappe.utils import cint, flt, getdate, now_datetime

from ms_calendar.api.field_merit_trac import _normalize_subject

MAX_ROWS = 20000
CHUNK = 1000

URL_DT = "Field Meritrac Test URL"
RES_DT = "Field MeritTrac Test Result"
FRF = "Field Registration Form"
ASSESS_DT = "Field Meritrac Assessment"
QPM = "Question Paper Master"
MARKS_DT = "Field Evaluator Marks"

# Question Paper Master: "40% & Above" selection, "25% to 39.50%", "Below 25%".
SELECT_PCT = 40
LOW_PCT = 25
BAND_SELECT = "40% & above"
BAND_MID = "25% – 39.5%"
BAND_LOW = "Below 25%"

# Registration Form statuses that mean "the test result hasn't been acted on".
PRE_TEST_STATUSES = {"", "New Applicant", "CV Shortlist", "Test Process"}

# start_time / end_time are free-text Data fields. The Initiate Test dialog
# saves "01 Sep 2026 11:24 PM"; older rows use "YYYY-MM-DD HH:MM:SS".
_WINDOW_FORMATS = (
    "%d %b %Y %I:%M %p",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
)


def _can_read(doctype):
    return frappe.db.exists("DocType", doctype) and frappe.has_permission(doctype, "read")


def _check_read():
    if not (_can_read(URL_DT) or _can_read(RES_DT)):
        frappe.throw("Not permitted to view MeritTrac test links or results", frappe.PermissionError)


def _str_list(value):
    """Client filter → clean list of strings (anything else is dropped)."""
    if not value:
        return []
    if isinstance(value, str):
        value = frappe.parse_json(value)
    if not isinstance(value, (list, tuple)):
        return []
    return [str(v).strip() for v in value if isinstance(v, (str, int)) and str(v).strip()]


def _parse_window(value):
    v = " ".join((value or "").split())
    for fmt in _WINDOW_FORMATS:
        try:
            return datetime.strptime(v, fmt)
        except ValueError:
            continue
    return None


def _http_url(value):
    """Only real http(s) links are passed to the page (never javascript: etc.)."""
    v = (value or "").strip()
    return v if v.lower().startswith(("https://", "http://")) else ""


def _score_pct(r):
    """Percentage for a result row: score_percentile is what merit_trac.py
    stores as the percentage; fall back to overall_percentage_score, then
    total_score / max_score."""
    for f in ("score_percentile", "overall_percentage_score"):
        if (r.get(f) or "") not in ("", None) and str(r.get(f)).strip().upper() != "NA":
            return round(flt(r.get(f)), 2)
    if flt(r.get("max_score")):
        return round(flt(r.get("total_score")) * 100 / flt(r.get("max_score")), 2)
    return None


def _band(pct):
    if pct is None:
        return ""
    if pct >= SELECT_PCT:
        return BAND_SELECT
    if pct >= LOW_PCT:
        return BAND_MID
    return BAND_LOW


def _chunks(values):
    values = list(values)
    for i in range(0, len(values), CHUNK):
        yield values[i:i + CHUNK]


def _get_list_in(doctype, field, values, fields, **kwargs):
    """frappe.get_list with an IN filter, chunked so huge IN lists stay sane."""
    out = []
    for part in _chunks(values):
        out += frappe.get_list(
            doctype, filters=[[doctype, field, "in", part]], fields=fields,
            limit_page_length=0, **kwargs
        )
    return out


def _qpm_index():
    """SA code → Question Paper Master rows (63 rows; one read)."""
    if not frappe.db.exists("DocType", QPM):
        return {}
    meta = frappe.get_meta(QPM)
    fields = [f for f in ("test_subject", "qp_hindi_english", "qp_kannada_english",
                          "test_result_selection_criteria", "evaluator_name", "evaluator_email")
              if f == "test_subject" or meta.has_field(f)]
    index = {}
    for r in frappe.get_all(QPM, fields=fields, limit_page_length=0):
        for code in (r.get("qp_hindi_english"), r.get("qp_kannada_english")):
            if (code or "").strip():
                index.setdefault(code.strip(), []).append(r)
    return index


def _qpm_row(index, sa_code, subject):
    """Same matching as field_merit_trac._get_evaluator_qp_row, in memory:
    the candidate's own subject decides when one SA code is shared."""
    rows = index.get((sa_code or "").strip()) or []
    if not rows:
        return None
    same = [r for r in rows if _normalize_subject(r.test_subject) == _normalize_subject(subject)]
    if same:
        return same[0]
    return rows[0] if len({(r.get("evaluator_email") or "") for r in rows}) == 1 else None


def _needs_evaluator(qp):
    if not qp:
        return False
    return bool((qp.get("evaluator_email") or "").strip()) or "evaluator" in (
        qp.get("test_result_selection_criteria") or ""
    ).lower()


# ── "Full data" tables: every field of both doctypes ─────────────────────
_SKIP_TYPES = {
    "Section Break", "Column Break", "Tab Break", "HTML", "Button", "Image",
    "Fold", "Heading", "Table", "Table MultiSelect",
}


def _raw_fields(doctype):
    """Every data field of the doctype this user may read (field-level
    permissions respected), in form order."""
    meta = frappe.get_meta(doctype)
    try:
        permitted = set(meta.get_permitted_fieldnames(permission_type="read"))
    except Exception:
        permitted = None
    out = []
    for df in meta.fields:
        if df.fieldtype in _SKIP_TYPES:
            continue
        if permitted is not None and df.fieldname not in permitted:
            continue
        if permitted is None and cint(df.permlevel):
            continue
        out.append({"fieldname": df.fieldname, "label": df.label or df.fieldname, "fieldtype": df.fieldtype})
    return out


def _with_fields(base, raw):
    seen, out = set(), []
    for f in base + [c["fieldname"] for c in raw]:
        if f not in seen:
            seen.add(f)
            out.append(f)
    return out


def _clean_text(value):
    return frappe.utils.strip_html(str(value or "")).strip()


@frappe.whitelist()
def get_dashboard_data(from_date=None, to_date=None, departments=None, roles=None,
                       subjects=None, statuses=None):
    """
    One row per applicant who was sent a MeritTrac link and/or has a result,
    restricted to what the user can read, plus the full records of both
    doctypes for the "Full data" tabs. Filter options come from all of the
    user's rows; cards / tables / drilldowns are derived on the page.
    """
    _check_read()
    return _build(from_date, to_date, departments, roles, subjects, statuses)


@frappe.whitelist(methods=["POST"])
def download_data(kind="both", from_date=None, to_date=None, departments=None, roles=None,
                  subjects=None, statuses=None):
    """
    Excel download of the "Full data" tables for the applied filters — the
    same permission-scoped _build the page uses, so the file can never hold
    more than the user sees. kind: "links", "results" or "both" (two sheets).
    """
    _check_read()
    if kind not in ("links", "results", "both"):
        frappe.throw("Unknown download")
    data = _build(from_date, to_date, departments, roles, subjects, statuses)

    def sheet(columns, records):
        return [[c["label"] for c in columns]] + [[r.get(c["fieldname"]) for c in columns] for r in records]

    sheets = []
    if kind in ("links", "both"):
        sheets.append(("Test URL", sheet(data["link_columns"], data["link_records"])))
    if kind in ("results", "both"):
        sheets.append(("Test Result", sheet(data["result_columns"], data["result_records"])))

    span = "-".join(filter(None, [data["from_date"], data["to_date"]])) or "all-dates"
    frappe.response["filename"] = f"field-merittrac-{kind}-{span}.xlsx"
    frappe.response["filecontent"] = _xlsx(sheets)
    frappe.response["type"] = "binary"


def _xlsx(sheets):
    """[(sheet name, rows)] → .xlsx bytes. openpyxl is on both v15 (cloud)
    and v16. Text is always written as text — a value starting with = + - @
    never becomes a formula — and characters Excel rejects are dropped."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for name, rows in sheets:
        ws = wb.create_sheet(name[:31])
        widths = {}
        for r_i, row in enumerate(rows, start=1):
            for c_i, value in enumerate(row, start=1):
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]
                cell = ws.cell(row=r_i, column=c_i, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if r_i == 1:
                    cell.font, cell.fill = head_font, head_fill
                widths[c_i] = max(widths.get(c_i, 8), min(60, len(str(value or "")) + 2))
        for c_i, w in widths.items():
            ws.column_dimensions[get_column_letter(c_i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build(from_date=None, to_date=None, departments=None, roles=None, subjects=None, statuses=None):
    now = now_datetime()
    from_date = getdate(from_date) if from_date else None
    to_date = getdate(to_date) if to_date else None
    if from_date and to_date and from_date > to_date:
        from_date, to_date = to_date, from_date

    departments = _str_list(departments)
    roles = _str_list(roles)
    subjects = _str_list(subjects)
    statuses = _str_list(statuses)

    # ── links sent ───────────────────────────────────────────────────────
    # (Each list also carries every other readable field, for "Full data".)
    links, link_raw = [], []
    if _can_read(URL_DT):
        link_raw = _raw_fields(URL_DT)
        links = frappe.get_list(
            URL_DT,
            fields=_with_fields(["name", "applicant_id", "applicant_name", "applicant_email", "applicant_role",
                                 "assessment_id", "subject", "start_time", "end_time", "creation", "owner"],
                                link_raw),
            order_by="creation asc",
            limit_page_length=MAX_ROWS,
        )

    # ── results received ─────────────────────────────────────────────────
    results, result_raw = [], []
    if _can_read(RES_DT):
        result_raw = _raw_fields(RES_DT)
        results = frappe.get_list(
            RES_DT,
            fields=_with_fields(["name", "applicant_id", "applicant_name", "attempt_status", "total_score",
                                 "max_score", "score_percentile", "overall_percentage_score", "credit_score",
                                 "total_questions", "total_attempted", "created_at", "score_report",
                                 "creation", "owner"],
                                result_raw),
            order_by="creation asc",
            limit_page_length=MAX_ROWS,
        )

    links_by_app, results_by_app = {}, {}
    for l in links:
        if l.applicant_id:
            links_by_app.setdefault(l.applicant_id, []).append(l)
    for r in results:
        if r.applicant_id:
            results_by_app.setdefault(r.applicant_id, []).append(r)
    applicants = set(links_by_app) | set(results_by_app)

    # ── applicant profile (only if the user can read Registration Forms) ─
    frf = {}
    if applicants and _can_read(FRF):
        for f in _get_list_in(FRF, "name", applicants,
                              ["name", "full_name_aadhaar", "email_address", "role", "department",
                               "written_subject", "application_status", "state"]):
            frf[f.name] = f

    # ── assessment set names for the SA codes on the links ──────────────
    assess = {}
    codes = {l.assessment_id for l in links if l.assessment_id}
    if codes and _can_read(ASSESS_DT):
        for a in _get_list_in(ASSESS_DT, "name", codes, ["name", "assessment_set"]):
            assess[a.name] = a.assessment_set or ""

    # ── evaluator marks, only for results this user can read ────────────
    marks = {}
    if results and _can_read(MARKS_DT):
        for m in _get_list_in(MARKS_DT, "test_result", [r.name for r in results],
                              ["name", "test_result", "result", "total_marks", "max_total_marks"]):
            marks[m.test_result] = m

    qpm = _qpm_index()

    rows = []
    for app in applicants:
        app_links = links_by_app.get(app, [])
        app_results = sorted(results_by_app.get(app, []),
                             key=lambda r: str(r.created_at or r.creation))
        f = frf.get(app) or {}
        # Current link = the one with the latest readable test window; if no
        # window can be read, the newest link record.
        dated = [(_parse_window(l.start_time), l) for l in app_links]
        dated = [x for x in dated if x[0]]
        link = (max(dated, key=lambda x: x[0])[1] if dated
                else app_links[-1] if app_links else None)
        res = app_results[-1] if app_results else None  # latest result

        start = _parse_window(link.start_time) if link else None
        end = _parse_window(link.end_time) if link else None

        if res:
            test_state = "Completed"
        elif not link:
            test_state = "No link"
        elif start and now < start:
            test_state = "Upcoming"
        elif end and now > end:
            test_state = "Not attempted"
        elif start or end:
            test_state = "Window open"
        else:
            test_state = "Window unknown"

        # Test date: when the result says the test was taken, else the window
        # start, else when the link record was made.
        test_dt = (res and (res.created_at or res.creation)) or start or (link and link.creation)
        test_date = getdate(test_dt) if test_dt else None

        subject = ((link.subject if link else "") or f.get("written_subject") or "").strip()
        sa_code = (link.assessment_id or "") if link else ""
        qp = _qpm_row(qpm, sa_code, subject)
        needs_eval = _needs_evaluator(qp)

        pct = _score_pct(res) if res else None
        m = marks.get(res.name) if res else None
        if not res or not needs_eval:
            eval_state = ""
        elif m:
            eval_state = "Marks submitted"
        else:
            eval_state = "Marks pending"

        status = (f.get("application_status") or "").strip()
        status_stale = bool(res) and status in PRE_TEST_STATUSES

        rows.append({
            "application_id": app,
            "applicant": f.get("full_name_aadhaar") or (link and link.applicant_name)
                or (res and res.applicant_name) or "",
            "email": f.get("email_address") or (link and link.applicant_email) or "",
            "role": (f.get("role") or (link and link.applicant_role) or "").strip(),
            "department": (f.get("department") or "").strip(),
            "state": f.get("state") or "",
            "subject": subject,
            "app_status": status,
            "status_stale": 1 if status_stale else 0,
            # links
            "links_sent": len(app_links),
            "link_name": link.name if link else "",
            "sa_code": sa_code,
            "assessment_set": assess.get(sa_code, ""),
            "window_start": start.strftime("%Y-%m-%d %H:%M") if start else "",
            "window_end": end.strftime("%Y-%m-%d %H:%M") if end else "",
            "window_label": (link.start_time or "") if link else "",
            "test_state": test_state,
            "test_date": str(test_date) if test_date else "",
            # result
            "result_name": res.name if res else "",
            "attempts": len(app_results),
            "score": flt(res.total_score) if res and (res.total_score or "") != "" else None,
            "max_score": flt(res.max_score) if res and (res.max_score or "") != "" else None,
            "pct": pct,
            "band": _band(pct),
            "dq_score": flt(res.credit_score) if res and (res.credit_score or "") not in ("", "NA") else None,
            "attempted": cint(res.total_attempted) if res and res.total_attempted else None,
            "questions": cint(res.total_questions) if res and res.total_questions else None,
            "attempt_status": (res.attempt_status or "") if res else "",
            "taken_on": str(res.created_at or "")[:16] if res else "",
            "report_url": _http_url(res.score_report) if res else "",
            # evaluator
            "needs_eval": 1 if needs_eval else 0,
            "evaluator": (qp.get("evaluator_name") or qp.get("evaluator_email") or "") if needs_eval else "",
            "eval_state": eval_state,
            "eval_result": (m.result or "") if m else "",
        })

    # Filter options from every row the user can read (before filtering).
    def distinct(key):
        return sorted({r[key] for r in rows if r[key]})

    options = {
        "departments": distinct("department"),
        "roles": distinct("role"),
        "subjects": distinct("subject"),
        "statuses": distinct("app_status"),
    }
    all_time = {
        "applicants": len(rows),
        "links": len(links),
        "results": len(results),
    }

    def in_range(d):
        if from_date and (not d or d < str(from_date)):
            return False
        if to_date and (not d or d > str(to_date)):
            return False
        return True

    def keep(r):
        return in_range(r["test_date"]) and keep_attrs(r)

    def keep_attrs(r):
        if departments and r["department"] not in departments:
            return False
        if roles and r["role"] not in roles:
            return False
        if subjects and r["subject"] not in subjects:
            return False
        if statuses and r["app_status"] not in statuses:
            return False
        return True

    out = sorted((r for r in rows if keep(r)),
                 key=lambda r: (r["test_date"], r["taken_on"]), reverse=True)

    # ── Full data: every record of both doctypes in the filters ─────────
    # Each record is dated by itself (a link by its test window start, a
    # result by when the test was taken); department / role / subject /
    # status come from its applicant.
    row_by_app = {r["application_id"]: r for r in rows}
    attr_filtered = bool(departments or roles or subjects or statuses)

    def record_ok(applicant_id, d):
        if not in_range(str(d) if d else ""):
            return False
        if attr_filtered:
            app_row = row_by_app.get(applicant_id)
            return bool(app_row) and keep_attrs(app_row)
        return True

    link_records = []
    for l in reversed(links):
        start = _parse_window(l.start_time)
        d = getdate(start or l.creation)
        if record_ok(l.applicant_id, d):
            rec = dict(l)
            rec["_test_date"] = str(d)
            link_records.append(rec)

    result_records = []
    for r in reversed(results):
        d = getdate(r.created_at or r.creation)
        if record_ok(r.applicant_id, d):
            rec = dict(r)
            rec["_test_date"] = str(d)
            result_records.append(rec)

    # Child rows (section scores / written answers) — only for results this
    # user can read, so no permission check is needed on the child tables.
    shown_results = {r["result_name"] for r in out if r["result_name"]}
    rec_by_name = {r["name"]: r for r in result_records}
    section_rows = []
    for part in _chunks(shown_results | set(rec_by_name)):
        section_rows += frappe.get_all(
            "Field MeritTrac Section Score",
            filters={"parenttype": RES_DT, "parent": ["in", part]},
            fields=["parent", "section_name", "score", "max_score", "idx"],
            order_by="idx asc",
            limit_page_length=0,
        )
    answers = {}
    for part in _chunks(rec_by_name):
        for a in frappe.get_all(
            "Field MeritTrac Descriptive Response",
            filters={"parenttype": RES_DT, "parent": ["in", part]},
            fields=["parent", "question_text", "candidate_response", "idx"],
            order_by="idx asc",
            limit_page_length=0,
        ):
            answers.setdefault(a.parent, []).append(a)
    for s in section_rows:
        rec = rec_by_name.get(s.parent)
        if rec is not None:
            rec["_sections"] = (rec.get("_sections", "") + "; " if rec.get("_sections") else "") + \
                f"{(s.section_name or '').strip()}: {flt(s.score):g} / {flt(s.max_score):g}"
    for name, rec in rec_by_name.items():
        qa = answers.get(name) or []
        rec["_answer_count"] = len(qa)
        rec["_answers"] = "\n\n".join(
            f"Q{i}. {_clean_text(a.question_text)}\nAnswer: {_clean_text(a.candidate_response) or '(not answered)'}"
            for i, a in enumerate(qa, start=1)
        )

    def columns(raw, extra=()):
        return ([{"fieldname": "_test_date", "label": "Test Date", "fieldtype": "Date"},
                 {"fieldname": "name", "label": "ID", "fieldtype": "Data"}]
                + raw + list(extra)
                + [{"fieldname": "creation", "label": "Created On", "fieldtype": "Datetime"},
                   {"fieldname": "owner", "label": "Created By", "fieldtype": "Data"}])

    link_columns = columns(link_raw) if link_raw else []
    result_columns = columns(result_raw, (
        {"fieldname": "_sections", "label": "Section-wise Scores", "fieldtype": "Small Text"},
        {"fieldname": "_answer_count", "label": "Written Answers (count)", "fieldtype": "Int"},
        {"fieldname": "_answers", "label": "Written Answers", "fieldtype": "Long Text"},
    )) if result_raw else []

    # Section-wise averages for the results still shown.
    sections = {}
    for s in section_rows:
        if s.parent not in shown_results:
            continue
        name = (s.section_name or "").strip() or "(unnamed)"
        agg = sections.setdefault(name, {"section": name, "count": 0, "score": 0.0, "max": 0.0, "results": []})
        agg["count"] += 1
        agg["score"] += flt(s.score)
        agg["max"] += flt(s.max_score)
        agg["results"].append(s.parent)

    return {
        "rows": out,
        "link_columns": link_columns,
        "link_records": link_records,
        "result_columns": result_columns,
        "result_records": result_records,
        "sections": sorted(sections.values(), key=lambda a: -a["count"]),
        "truncated": len(links) >= MAX_ROWS or len(results) >= MAX_ROWS,
        "from_date": str(from_date) if from_date else "",
        "to_date": str(to_date) if to_date else "",
        "now": now.strftime("%Y-%m-%d %H:%M"),
        "options": options,
        # All-time cards: ignore every filter, still permission-scoped.
        "all_time": all_time,
        "thresholds": {"select": SELECT_PCT, "low": LOW_PCT},
    }


@frappe.whitelist()
def get_candidate_detail(application_id):
    """
    Second-level drilldown: everything about one applicant's MeritTrac test —
    profile, every link sent, every result in full (section scores, written
    answers, proctoring notes) and the evaluator's marks. Checks read
    permission on this exact Field Registration Form record first.
    """
    _check_read()
    application_id = (application_id or "").strip()
    if not application_id or not frappe.db.exists(FRF, application_id):
        frappe.throw("Applicant not found")
    if not frappe.has_permission(FRF, "read", doc=application_id):
        frappe.throw("Not permitted to view this applicant", frappe.PermissionError)

    profile_fields = [
        "full_name_aadhaar", "email_address", "phone_number", "role", "department",
        "job_code", "location", "state", "district_city", "written_subject",
        "preferred_test_mode", "test_location", "application_status", "application_date",
    ]
    meta = frappe.get_meta(FRF)
    profile_fields = [f for f in profile_fields if meta.has_field(f)]
    pdf_field = ["filed_merit_track_test"] if meta.has_field("filed_merit_track_test") else []
    frf = frappe.db.get_value(FRF, application_id, profile_fields + pdf_field, as_dict=True) or {}
    # Short labels where the form's own label is a whole question.
    short_labels = {"test_location": "Test location", "preferred_test_mode": "Preferred test mode"}

    def fmt(f):
        value = frf.get(f)
        if value in (None, ""):
            return ""
        if meta.get_field(f).fieldtype in ("Date", "Datetime"):
            return frappe.utils.formatdate(value)
        return str(value)

    profile = [{"label": short_labels.get(f) or meta.get_label(f), "value": fmt(f)} for f in profile_fields]
    pdf = frf.get("filed_merit_track_test") or ""
    pdf = pdf if pdf.startswith(("/files/", "/private/files/")) else ""

    links = []
    if _can_read(URL_DT):
        assess_names = {}
        rows = frappe.get_list(
            URL_DT, filters={"applicant_id": application_id},
            fields=["name", "assessment_id", "subject", "start_time", "end_time", "creation"],
            order_by="creation asc", limit_page_length=100,
        )
        codes = {r.assessment_id for r in rows if r.assessment_id}
        if codes and _can_read(ASSESS_DT):
            assess_names = {a.name: a.assessment_set for a in frappe.get_list(
                ASSESS_DT, filters={"name": ["in", list(codes)]}, fields=["name", "assessment_set"])}
        now = now_datetime()
        for r in rows:
            s, e = _parse_window(r.start_time), _parse_window(r.end_time)
            state = ("Upcoming" if s and now < s else "Closed" if e and now > e
                     else "Open" if (s or e) else "Unknown")
            links.append({
                "name": r.name, "sa_code": r.assessment_id or "",
                "assessment_set": assess_names.get(r.assessment_id, ""),
                "subject": r.subject or "", "start": r.start_time or "", "end": r.end_time or "",
                "state": state, "sent_on": str(r.creation)[:16],
            })

    results = []
    if _can_read(RES_DT):
        for name in frappe.get_list(RES_DT, filters={"applicant_id": application_id}, pluck="name",
                                    order_by="creation asc", limit_page_length=50):
            doc = frappe.get_doc(RES_DT, name)
            pct = _score_pct(doc)
            mark = None
            if _can_read(MARKS_DT):
                m = frappe.get_list(MARKS_DT, filters={"test_result": name},
                                    fields=["name", "evaluator_name", "total_marks", "max_total_marks",
                                            "result", "comments", "creation"],
                                    limit_page_length=1)
                if m:
                    m = m[0]
                    mark = {
                        "name": m.name, "evaluator": m.evaluator_name or "",
                        "total": m.total_marks, "max": m.max_total_marks, "result": m.result or "",
                        "comments": frappe.utils.strip_html(m.comments or "").strip(),
                        "on": str(m.creation)[:16],
                    }
            results.append({
                "name": doc.name,
                "taken_on": str(doc.created_at or "")[:16],
                "attempt_status": doc.attempt_status or "",
                "score": doc.total_score or "", "max_score": doc.max_score or "",
                "pct": pct, "band": _band(pct),
                "attempted": doc.total_attempted or "", "questions": doc.total_questions or "",
                "dq_score": doc.get("credit_score") or "",
                "proctor_comment": frappe.utils.strip_html(doc.get("proctor_comment") or "").strip(),
                "report_url": _http_url(doc.score_report),
                "tn_report_url": _http_url(doc.get("tn_report")),
                "sections": [
                    {"section": s.section_name or "", "score": flt(s.score), "max": flt(s.max_score)}
                    for s in (doc.get("section_wise_scores") or [])
                ],
                # Answers can be HTML from MeritTrac's editor — sent as plain
                # text, the page escapes it.
                "answers": [
                    {"q": frappe.utils.strip_html(a.question_text or "").strip(),
                     "a": frappe.utils.strip_html(a.candidate_response or "").strip()}
                    for a in (doc.get("descriptive_responses") or [])
                ],
                "marks": mark,
            })

    return {
        "application_id": application_id,
        "name": frf.get("full_name_aadhaar") or application_id,
        "profile": profile,
        "pdf": pdf,
        "links": links,
        "results": results,
        "thresholds": {"select": SELECT_PCT, "low": LOW_PCT},
    }
