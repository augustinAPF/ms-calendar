import io
import json
import re
from datetime import timedelta

import frappe
from frappe.utils import getdate, today


def _parse_multi(val):
	"""Turns the native_state/school querystring param into a list of
	selected values. reports.js's State/School filters are now chip
	multiselects (see makeMultiselect) and send their selection
	JSON-encoded — not comma-joined, since school names like "Khargone,
	Madhya Pradesh" already contain commas and a plain join/split would
	wrongly split one school into two. Falls back to treating `val` as a
	single plain value if it isn't valid JSON, so a hand-typed URL (or any
	older caller) with a bare string still works as a one-item filter.
	"""
	if not val:
		return []
	try:
		parsed = json.loads(val)
	except (TypeError, ValueError):
		return [val]
	if isinstance(parsed, list):
		return [str(v).strip() for v in parsed if str(v).strip()]
	return [str(parsed).strip()] if str(parsed).strip() else []


# Same role/column mapping as the "School Teacher — District Funnel" report
# in ms_calendar/ms_calendar/page/reports/reports.js (STF_COLS there) — keep
# these two in sync if either changes. Duplicated rather than shared because
# one side is JS running in the browser and this side is Python running on
# the server; there's no natural single source of truth to point both at
# without a bigger refactor of that page.
#
# application_status on Field Registration Form carries two overlapping
# naming eras for the same stage at once (confirmed 2026-09-04) — every
# column below lists every real spelling variant, not just one.
STF_COLS = [
	("No. of Applications", "applications", None),
	("CV Shortlisted", "cv_shortlisted", ["CV Shortlist", "Shortlisted"]),
	("CV Screening Pending", "cv_pending", ["New Applicant", "Pending", "Correction Pending", "Pending for Correction", "PI Edited"]),
	("CV Rejected", "cv_rejected", ["CV Reject", "Rejected", "Not Selected", "Duplicated", "Blacklisted", "Blocklisted", "Blocklist"]),
	("On Hold", "on_hold", ["On Hold", "Test Hold"]),
	("Admit Card Sent", "admit_card_sent", ["Test Process", "Admit Card Sent"]),
	("Test Select", "assessment_pass", ["Test Select", "Assessment Passed"]),
	("Test Reject", "assessment_fail", ["Test Reject", "Assessment Failed"]),
	("Recruiter Round Scheduled", "rr_sched", ["Recruiter Round", "Recruiter Round Scheduled", "Recruiter Round-Interview Scheduled"]),
	("Recruiter Round Selected", "rr_sel", ["Recruiter Round Select", "Recruiter Round-Interview Selected"]),
	("Recruiter Round Rejected", "rr_rej", ["Recruiter Round Reject", "Recruiter Reject", "Recruiter Round-Interview Rejected"]),
	("Subject Round Interview Scheduled", "sr_sched",
		["Functional Round", "Education Capacity Round", "Round One", "Subject Round", "Functional Round Scheduled", "Education Capacity Round Scheduled", "Subject Round Scheduled", "Subject Round-Interview Scheduled",
			"Functional Round-Interview Scheduled", "Functional Round-Interviewed"]),
	("Subject Round Interview Selected", "sr_sel",
		["Functional Round Select", "Education Capacity Select", "Round One Select", "Subject Round Select", "Functional Round-Interview Selected"]),
	("Subject Round Interview Rejected", "sr_rej",
		["Functional Round Reject", "Education Capacity Reject", "Round 1 Reject", "Subject Round Reject", "Subject Round-Interview Rejected",
			"Functional Round-Interview Rejected"]),
	("Demo Round Scheduled", "dr_sched",
		["Round Two", "Demo Round", "Demo Round Scheduled", "Subject/Classroom Demo Round-Interview Scheduled", "Principal/Demo Round-Interview Scheduled"]),
	("Demo Round Selected", "dr_sel",
		["Round Two Select", "Demo Round Select", "Subject/Classroom Demo Round-Interview Selected", "Principal/Demo Round-Interview Selected"]),
	("Demo Round Rejected", "dr_rej",
		["Round 2 Reject", "Demo Round Reject", "Subject/Classroom Demo Round-Interview Rejected", "Principal/Demo Round-Interview Rejected",
			"Demo Round-Interview Rejected", "Demo & Leader Round-Interview Rejected"]),
	("Leader Round 1 Interview Scheduled", "lr1_sched", ["Leader Round-1", "Leader Round-1 Scheduled", "Leader Round 1-Interview Scheduled", "Leader Round 1-Interviewed"]),
	("Leader Round 1 Interview Selected", "lr1_sel", ["Leader Round-1 Select", "Leader Round 1-Interview Selected", "Leader Round-Interview Selected"]),
	("Leader Round 1 Interview Rejected", "lr1_rej", ["Leader Round-1 Reject", "Leader Round 1-Interview Rejected"]),
	("Leader Round 2 Interview Scheduled", "lr2_sched", ["Leader Round-2", "Leader Round-2 Scheduled"]),
	("Leader Round 2 Interview Selected", "lr2_sel", ["Leader Round-2 Select"]),
	("Leader Round 2 Interview Rejected", "lr2_rej", ["Leader Round-2 Reject"]),
	("Leader Round 3 Interview Scheduled", "lr3_sched", ["Round Three", "Leader Round-3", "Leader Round-3 Scheduled", "Leader Round 3-Interview Scheduled", "Leader Round 3-Interviewed"]),
	("Leader Round 3 Interview Selected", "lr3_sel", ["Round Three Select", "Leader Round-3 Select", "Leader Round 3-Interview Selected"]),
	("Leader Round 3 Interview Rejected", "lr3_rej", ["Round 3 Reject", "Leader Round-3 Reject", "Leader Round 3-Interview Rejected"]),
	("Pending with CBT (Offer Released)", "pending_cbt", ["Pending With CBT", "CBT Assigned", "Offer", "Offer Sent"]),
]


def _get_col(role):
	"""Same substring match as reports.js's getCol() — real role values are
	region-suffixed ("School Teacher - Barmer"), an exact match only ever
	caught the plain unsuffixed value."""
	r = (role or "").strip().lower()
	if "teacher" in r:
		return "ST"
	if "health" in r or "livelihood" in r:
		return None
	return "RP"


def _detail_fields():
	"""Every real, per-applicant field on Field Registration Form — layout
	fields (Section/Column/Tab Break, HTML, Button) and child tables
	excluded, since those aren't values that belong in one flat row.
	Built from the doctype's own meta rather than a hand-maintained list,
	so a field added/renamed there shows up in the Details sheet
	automatically without this file needing a matching edit — "whole data"
	is meant literally here (2026-09-07 request), not a curated subset.
	Returns [(fieldname, label), ...] in form order; a fieldname repeated
	more than once in the doctype (confirmed: "zayam_id" is on this one)
	is only included the first time.
	"""
	skip_types = {
		"Section Break", "Column Break", "Tab Break", "HTML", "Button", "Fold",
		"Table", "Table MultiSelect",
	}
	meta = frappe.get_meta("Field Registration Form")
	seen = set()
	out = []
	for f in meta.fields:
		if f.fieldtype in skip_types or f.fieldname in seen:
			continue
		seen.add(f.fieldname)
		out.append((f.fieldname, f.label or f.fieldname))
	return out


def _district_of(rec):
	""""School" is derived from the Role field's own text, not
	native_district/location (confirmed 2026-09-04) — see the matching
	districtOf() in reports.js for the full explanation; kept identical
	here so the Excel download and the on-screen table never disagree."""
	role = (rec.get("role") or "").strip()
	if not role:
		return "Unspecified"
	dash_idx = role.rfind(" - ")
	if dash_idx >= 0:
		return role[dash_idx + 3:].strip() or "Unspecified"
	return role


# The report only ever covers these 12 known schools (confirmed
# 2026-09-11) — same real districtOf() values as FIXED_SCHOOLS in
# reports.js (kept in sync manually, like STF_COLS above). Every other
# distinct _district_of() bucket in the data (unsuffixed roles like
# "School Teacher Education", "Resource Person / Teacher Educator",
# "Azim Premji School: ...", state-only buckets like "Rajasthan", etc.)
# is noise nobody wants in this report, on screen or in the download.
STF_SCHOOL_VALUES = [
	"Bengaluru", "Kalaburagi", "Yadgir", "Dhamtari", "Raigarh", "Barmer",
	"Sirohi", "Tonk", "Itki, Ranchi", "Udham Singh Nagar", "Uttarkashi",
	"Khargone, Madhya Pradesh",
]


@frappe.whitelist()
def download_st_district_funnel(
	applied_from=None, applied_to=None, native_state=None, school=None, job_code=None
):
	"""Server-rendered .xlsx of the "School Teacher — District Funnel"
	report, with the same header/row background colours as the on-screen
	table (reports.js) — the CSV export used elsewhere on this page has no
	concept of colour, so this needed a real spreadsheet library instead
	(openpyxl, already used for exports elsewhere in this session).
	"""
	filters = {}
	# Optional range on application_date (the form's "Application Date" — a
	# Datetime field, so the To day runs to 23:59:59) — the only date filter,
	# same as reports.js's Application Date From/To (the creation From/To
	# range was removed and date_of_applied swapped for this, 2026-10-15).
	if applied_from or applied_to:
		filters["application_date"] = [
			"between",
			[(applied_from or "2000-01-01") + " 00:00:00", (applied_to or today()) + " 23:59:59"],
		]

	# "name" and "creation" are the two fields the Details sheet needs that
	# aren't in _detail_fields() (they're the docname/ID and a standard
	# system field, not real meta fields on the doctype) — everything else
	# comes straight from the doctype's own field list, so the Details
	# sheet always has "whole data", not a hand-picked subset.
	detail_fields = _detail_fields()
	rows = frappe.get_all(
		"Field Registration Form",
		filters=filters,
		fields=["name", "creation"] + [fn for fn, _label in detail_fields],
		limit_page_length=0,
	)
	st_rows = [r for r in rows if _get_col(r.role) == "ST"]

	# Restricted to only the 12 known schools (see STF_SCHOOL_VALUES) —
	# same restriction reports.js applies before the table ever renders,
	# so the download always matches the on-screen report even with no
	# School filter selected.
	st_rows = [r for r in st_rows if _district_of(r) in STF_SCHOOL_VALUES]

	# Same State/School/Job Code multiselect filters as the on-screen table
	# (reports.js's applySTFFilters) — passed through as querystring params
	# from the Download Excel button so the file matches whatever's
	# currently filtered on screen, instead of always exporting every
	# School Teacher record in the date range. Each is zero or more
	# selected values (see _parse_multi) — empty means "All" (no
	# filtering), a non-empty list matches any of the selected values.
	states = _parse_multi(native_state)
	schools = _parse_multi(school)
	job_codes = _parse_multi(job_code)
	if states:
		st_rows = [r for r in st_rows if (r.get("native_state") or "").strip() in states]
	if schools:
		st_rows = [r for r in st_rows if _district_of(r) in schools]
	if job_codes:
		st_rows = [r for r in st_rows if (r.get("job_code") or "").strip() in job_codes]

	districts = sorted({_district_of(r) for r in st_rows})

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.hyperlink import Hyperlink

	wb = Workbook()
	ws = wb.active
	ws.title = "School Teacher Funnel"

	# Same palette as the on-screen table in reports.js's showSTDistrictFunnel.
	corner_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
	corner_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill(start_color="F5A623", end_color="F5A623", fill_type="solid")
	header_font = Font(bold=True, color="3B2C00")
	school_fill = PatternFill(start_color="FFF3D6", end_color="FFF3D6", fill_type="solid")
	total_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
	total_font = Font(bold=True)
	thin_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

	corner = ws.cell(row=1, column=1, value="School")
	corner.fill = corner_fill
	corner.font = corner_font
	corner.alignment = thin_center

	for ci, (label, _key, _statuses) in enumerate(STF_COLS, start=2):
		c = ws.cell(row=1, column=ci, value=label)
		c.fill = header_fill
		c.font = header_font
		c.alignment = thin_center

	# "Details" sheet — the Excel equivalent of clicking a number cell on
	# the on-screen report (reports.js's stf-cell click handler, which
	# opens a dialog of the matching records). Excel cells can't pop up a
	# dialog on click, so instead every non-zero count cell on the summary
	# sheet becomes a hyperlink that jumps straight to that exact block of
	# matching records here, and the sheet has AutoFilter turned on so it
	# can be narrowed further by hand. Built in the same School x Stage
	# loop as the summary counts below, so a block's row range and its
	# cell's hyperlink target can never drift out of sync.
	hyperlink_font = Font(color="0563C1", underline="single")
	details_ws = wb.create_sheet("Details")
	# School/Stage (this report's own grouping, not on the doctype) +
	# Applicant ID (the docname) up front, then literally every field
	# _detail_fields() found on Field Registration Form, in form order —
	# "whole data" per applicant, not a hand-picked subset.
	details_headers = ["School", "Stage", "Applicant ID"] + [label for _fn, label in detail_fields] + ["Date"]
	last_col_letter = get_column_letter(len(details_headers))

	# Row 1: a plain-language instruction banner, not just column headers.
	# A hyperlink can jump here but can't apply a filter on its own (no
	# code runs on click in a static .xlsx, and OnlyOffice doesn't
	# reliably support the macro event that would need). The AutoFilter
	# buttons on the School/Stage columns below already do the "show only
	# this, hide everything else" job by hand in two clicks — this banner
	# spells that out right in the file instead of assuming it's obvious.
	details_ws.merge_cells(f"A1:{last_col_letter}1")
	banner = details_ws.cell(
		row=1, column=1,
		value=(
			"To see ONLY one group and hide everything else: click the ▼ on the "
			"School column below, tick just that School — then click the ▼ on "
			"Stage and tick just that Stage. Clear both filters (▼ → Clear Filter) "
			"to see everything again."
		),
	)
	banner.font = Font(bold=True, color="7D4F00")
	banner.fill = PatternFill(start_color="FFF3D6", end_color="FFF3D6", fill_type="solid")
	banner.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
	details_ws.row_dimensions[1].height = 30

	for ci, h in enumerate(details_headers, start=1):
		c = details_ws.cell(row=2, column=ci, value=h)
		c.fill = header_fill
		c.font = header_font
	details_ws.freeze_panes = "A3"
	details_ws.auto_filter.ref = f"A2:{last_col_letter}2"
	details_row = 3  # next free row on the Details sheet

	# A real click-driven filter isn't possible in a static .xlsx (that
	# needs a macro, and OnlyOffice doesn't reliably support the
	# FollowHyperlink event a macro would need). Every hyperlink already
	# jumps to a block that is ONLY that School+Stage's matching rows,
	# nothing else mixed in — this banding just makes that block boundary
	# obvious at a glance instead of requiring the reader to notice the
	# School/Stage columns changed.
	# Both tints, not white — a white block is indistinguishable from the
	# sheet's own blank background (confirmed 2026-09-07: it just looked
	# unshaded), which defeated the point of banding every other block.
	# Same pastel pair as STATE_COLORS in reports.js, so this reads as the
	# same visual language as the rest of the report.
	block_fills = [
		PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
		PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
	]
	block_top = Side(style="medium", color="404040")
	block_bottom = Side(style="medium", color="404040")
	block_side = Side(style="thin", color="BFBFBF")
	block_parity = 0  # alternates 0/1 so consecutive blocks never share a fill

	grand = {key: 0 for (_label, key, _statuses) in STF_COLS}
	row_idx = 2
	for d in districts:
		recs = [r for r in st_rows if _district_of(r) == d]
		name_cell = ws.cell(row=row_idx, column=1, value=d)
		name_cell.fill = school_fill
		name_cell.font = Font(bold=True)

		for ci, (label, key, statuses) in enumerate(STF_COLS, start=2):
			if statuses is None:
				matched = recs
			else:
				matched = [r for r in recs if (r.application_status or "").strip() in statuses]
			cnt = len(matched)
			grand[key] += cnt

			cell = ws.cell(row=row_idx, column=ci, value=cnt or None)
			if matched:
				block_start = details_row
				block_end = details_row + len(matched) - 1
				fill = block_fills[block_parity]
				block_parity = 1 - block_parity
				for i, r in enumerate(matched):
					is_first = (details_row == block_start)
					is_last = (details_row == block_end)
					border = Border(
						top=block_top if is_first else None,
						bottom=block_bottom if is_last else None,
						left=block_side, right=block_side,
					)
					row_values = (
						[d, label, r.name]
						+ [r.get(fn) for fn, _label in detail_fields]
						+ [str(r.creation).split(" ")[0] if r.creation else ""]
					)
					for col, val in enumerate(row_values, start=1):
						dc = details_ws.cell(row=details_row, column=col, value=val)
						dc.fill = fill
						dc.border = border
					details_row += 1
				# location= (not a "#..." string on cell.hyperlink directly)
				# is what makes openpyxl write a plain internal `location`
				# attribute with no relationship at all. Assigning a bare
				# "#'Details'!A1" string instead makes openpyxl create an
				# actual external hyperlink relationship whose Target
				# literally is that "#..." text with TargetMode="External"
				# — which is exactly what was making OnlyOffice (and Excel)
				# show a "this link may be unsafe" warning on every click,
				# despite the link only ever pointing inside this same file.
				cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'Details'!A{block_start}")
				cell.font = hyperlink_font
		row_idx += 1

	total_name = ws.cell(row=row_idx, column=1, value="Total")
	total_name.fill = total_fill
	total_name.font = total_font
	for ci, (_label, key, _statuses) in enumerate(STF_COLS, start=2):
		c = ws.cell(row=row_idx, column=ci, value=grand[key] or None)
		c.fill = total_fill
		c.font = total_font

	# details_headers now runs to ~55 columns ("whole data" per applicant,
	# not a hand-picked subset — see _detail_fields()), so fixed per-letter
	# widths aren't practical to hand-maintain. Size each one off its own
	# header text length instead — short ones (Age, DOB) stay narrow, long
	# ones (the various "If Yes, please mention..." labels) get more room,
	# capped so one very long label can't blow out the whole sheet.
	for ci, h in enumerate(details_headers, start=1):
		details_ws.column_dimensions[get_column_letter(ci)].width = min(40, max(12, len(h) + 4))

	ws.freeze_panes = "B2"
	ws.column_dimensions["A"].width = 22
	for ci in range(2, len(STF_COLS) + 2):
		ws.column_dimensions[get_column_letter(ci)].width = 15
	ws.row_dimensions[1].height = 45

	buf = io.BytesIO()
	wb.save(buf)
	buf.seek(0)

	frappe.local.response.filename = "School_Teacher_District_Funnel.xlsx"
	frappe.local.response.filecontent = buf.getvalue()
	frappe.local.response.type = "binary"




# ── Job-opening state ─────────────────────────────────────────────────────
# Every report's "State" is the state of the JOB the candidate applied for,
# never the candidate's own native_state (recruiter feedback 2026-10-27:
# "West Bengal" showed as a pipeline column because a candidate lived
# there). Source order: the job location saved on the application
# (Field Registration Form.location — already narrowed to one option when
# the opening offered several), then the linked Job Opening's location.
# Both are free text ("Itki, Ranchi, Jharkhand, India", "Raigarh -
# Chhattisgarh/Gumla - Jharkhand", ...), so they're matched against state
# names first, then city/district names that show up without their state.
# reports.js / field_over_all_dashb.js get these lists from get_report_meta()
# and run the same matching (jobStateOf), so all reports agree.
INDIA_STATES = [
	"Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar", "Chhattisgarh", "Goa", "Gujarat",
	"Haryana", "Himachal Pradesh", "Jharkhand", "Karnataka", "Kerala", "Madhya Pradesh",
	"Maharashtra", "Manipur", "Meghalaya", "Mizoram", "Nagaland", "Odisha", "Punjab", "Rajasthan",
	"Sikkim", "Tamil Nadu", "Telangana", "Tripura", "Uttar Pradesh", "Uttarakhand", "West Bengal",
	"Andaman and Nicobar Islands", "Chandigarh", "Dadra and Nagar Haveli and Daman and Diu", "Delhi",
	"Jammu and Kashmir", "Ladakh", "Lakshadweep", "Puducherry",
]
STATE_ALIASES = {
	"Chhattisgarh": ["chattisgarh", "chhatisgarh", "dhamtari", "raigarh", "raipur", "bilaspur", "ambikapur"],
	"Karnataka": ["bengaluru", "bangalore", "kalaburagi", "yadgir", "jevargi"],
	"Madhya Pradesh": ["khargone", "bhopal", "damoh", "sagar"],
	"Rajasthan": ["tonk", "sirohi", "barmer", "udaipur", "banswara", "chittorgarh"],
	"Uttarakhand": ["uttaranchal", "udham singh", "uttarkashi", "dehradun"],
	"Jharkhand": ["ranchi", "dhanbad", "jamshedpur", "deoghar", "pakur", "sahibganj", "gumla", "simdega", "godda"],
	"Uttar Pradesh": ["lucknow", "varanasi", "ayodhya"],
	"Odisha": ["orissa", "rayagada", "keonjhar", "sambalpur"],
	"Assam": ["guwahati", "jorhat"],
	"Tamil Nadu": ["tamilnadu", "chennai"],
	"Telangana": ["hyderabad"],
	"West Bengal": ["kolkata"],
	"Puducherry": ["pondicherry"],
	"Maharashtra": ["nagpur", "mumbai", "pune"],
	"Andaman and Nicobar Islands": ["andaman"],
	"Jammu and Kashmir": ["jammu", "kashmir"],
}
STATE_MULTI = "Multiple States"
STATE_NONE = "Other / Not specified"


def _states_in(text):
	s = (text or "").lower()
	if not s:
		return []
	found = [st for st in INDIA_STATES if st.lower() in s]
	if not found:
		found = [st for st, keys in STATE_ALIASES.items() if any(k in s for k in keys)]
	return found


def job_state(location, jo_location=None):
	multi = False
	for text in (location, jo_location):
		found = _states_in(text)
		if len(found) == 1:
			return found[0]
		if found or "all states" in (text or "").lower():
			multi = True
	return STATE_MULTI if multi else STATE_NONE


def _job_opening_locations():
	return {r.name: r.location for r in frappe.get_all("Job Opening", fields=["name", "location"])}


# ── Workflow rounds & pipeline rows ───────────────────────────────────────
# Interview rounds in workflow order. Names are the Field Registration Form
# Application Status values: "<round>" = in that round, then its select /
# reject status. Rank = position in the hiring flow; rounds sharing a rank
# are alternatives (a role has one of EC / Subject / Technical /
# Functional). Used to tell whether a candidate got past a round.
WORKFLOW_ROUNDS = [
	("Recruiter Round", "Recruiter Round Select", "Recruiter Round Reject", 1),
	("Education Capacity Round", "Education Capacity Select", "Education Capacity Reject", 2),
	("Subject Round", "Subject Round Select", "Subject Round Reject", 2),
	("Technical Round", "Technical Round Select", "Technical Round Reject", 2),
	("Functional Round", "Functional Round Select", "Functional Round Reject", 2),
	("Demo Round", "Demo Round Select", "Demo Round Reject", 3),
	("Leader Round-1", "Leader Round-1 Select", "Leader Round-1 Reject", 4),
	("Leader Round-2", "Leader Round-2 Select", "Leader Round-2 Reject", 5),
	("Leader Round-3", "Leader Round-3 Select", "Leader Round-3 Reject", 6),
]
ROUND_RANK = {r[0]: r[3] for r in WORKFLOW_ROUNDS}
MAX_ROUND_RANK = 6
# Older status / interview-round spellings (same mapping as STF_COLS above).
_LEGACY_ROUNDS = [
	("recruiter", "Recruiter Round"),
	("round one", "Functional Round"), ("round 1", "Functional Round"),
	("round two", "Demo Round"), ("round 2", "Demo Round"),
	("round three", "Leader Round-3"), ("round 3", "Leader Round-3"),
]
# Statuses that come after every interview round.
_AFTER_ROUNDS = ("calibration", "document", "fitment", "compensation", "offer", "addendum",
	"pre joining", "it team", "it email", "joined", "erp id")


def canon_round(text):
	"""Workflow round name for an Application Status / Interview Round
	value ("Leader Round-2 Reject" -> "Leader Round-2"), or None."""
	s = re.sub(r"leader round (\d)", r"leader round-\1", (text or "").strip().lower())
	if not s:
		return None
	for name, sel, rej, _rank in WORKFLOW_ROUNDS:
		if s.startswith(name.lower()) or s in (sel.lower(), rej.lower()):
			return name
	for prefix, name in _LEGACY_ROUNDS:
		if s.startswith(prefix):
			return name
	return None


def _round_result(status):
	"""'select' / 'reject' / 'in' for a round status."""
	s = (status or "").lower()
	if "select" in s:
		return "select"
	if "reject" in s:
		return "reject"
	return "in"


def passed_rank(status):
	"""Highest round rank the candidate's current status shows they passed
	(0 = none / unknown, e.g. Not Selected, Interview No Show)."""
	rnd = canon_round(status)
	if rnd:
		return ROUND_RANK[rnd] - (0 if _round_result(status) == "select" else 1)
	s = (status or "").lower()
	if any(k in s for k in _AFTER_ROUNDS):
		return MAX_ROUND_RANK
	return 0


def _slug(text):
	return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


_STATUS_ROW = {
	"New Applicant": "cv_pending", "Applied": "cv_pending",
	"CV Shortlist": "cv_shortlist", "Shortlisted": "cv_shortlist",
	"CV Reject": "cv_reject", "Rejected": "cv_reject",
	"On Hold": "cv_hold",
	"Test Initiated": "test_process", "Test Process": "test_process",
	"Test Select": "test_select", "Test Reject": "test_reject", "Test Hold": "test_hold",
	"No Show": "test_no_show", "Test No Show": "test_no_show",
	"Interview No Show": "interview_no_show", "Interview-No Show": "interview_no_show",
	"Calibration Process": "calibration_process", "Calibration Select": "calibration_select",
	"Calibration Reject": "calibration_reject",
	"Not Selected": "not_selected", "Blocklisted": "blocklisted", "Blocklist": "blocklisted",
	"Duplicated": "duplicated", "Pending for Correction": "correction_pending",
	"Correction Pending": "correction_pending", "PI Edited": "pi_edited",
	"Selected": "o_joined",
}
_OFFER_ROWS = [
	"Document Collection", "Document Verification Pass", "Document Verification Fail",
	"Fitment Stage", "Fitment Reject", "Compensation Due", "Compensation Approved",
	"Offer", "Offer Sent", "Offer Not Sent", "Offer Accepted", "Offer Declined", "Offer Revoked",
	"Addendum Accepted", "Addendum Declined", "Pre joining", "IT Team", "IT Email Sent",
	"IT Email Not Sent", "Joined", "Not Joined", "ERP ID Created", "ERP ID Not Created",
]


def pipeline_stages():
	"""Pipeline report rows, in workflow order: [{stage, rows: [{label, key,
	total?}]}]. Each interview round has Select / Reject plus three rows for
	candidates still in it: Scheduled (interview date today or later),
	Feedback Pending (interview held, no result yet) and Pending (no
	interview scheduled yet) — see pipeline_row_key()."""
	stages = [
		{"stage": "Applications", "rows": [
			{"label": "Carried forward (applied before 1 April this FY)", "key": "carried_forward"},
			{"label": "Received from 1 April this FY", "key": "received_this_year"},
			{"label": "Total Applications", "key": "total", "total": True},
		]},
		{"stage": "CV Screening", "rows": [
			{"label": "CV Shortlist", "key": "cv_shortlist"},
			{"label": "CV Reject", "key": "cv_reject"},
			{"label": "New Applicant (Pending)", "key": "cv_pending"},
			{"label": "On Hold", "key": "cv_hold"},
		]},
		{"stage": "Written Test", "rows": [
			{"label": "Test Initiated / In Process", "key": "test_process"},
			{"label": "Test Select", "key": "test_select"},
			{"label": "Test Reject", "key": "test_reject"},
			{"label": "Test Hold", "key": "test_hold"},
			{"label": "No Show", "key": "test_no_show"},
		]},
	]
	for name, sel, rej, _rank in WORKFLOW_ROUNDS:
		k = _slug(name)
		stages.append({"stage": name, "rows": [
			{"label": sel, "key": k + "_select"},
			{"label": rej, "key": k + "_reject"},
			{"label": "Scheduled", "key": k + "_scheduled"},
			{"label": "Feedback Pending", "key": k + "_feedback"},
			{"label": "Pending (not scheduled)", "key": k + "_pending"},
		]})
	stages += [
		{"stage": "Interview No Show", "rows": [{"label": "Interview No Show", "key": "interview_no_show"}]},
		{"stage": "Calibration", "rows": [
			{"label": "Calibration Process", "key": "calibration_process"},
			{"label": "Calibration Select", "key": "calibration_select"},
			{"label": "Calibration Reject", "key": "calibration_reject"},
		]},
		{"stage": "Offer & Joining", "rows": [{"label": s, "key": "o_" + _slug(s)} for s in _OFFER_ROWS]},
		{"stage": "Closed / Other", "rows": [
			{"label": "Not Selected", "key": "not_selected"},
			{"label": "Blocklisted", "key": "blocklisted"},
			{"label": "Duplicated", "key": "duplicated"},
			{"label": "Correction Pending", "key": "correction_pending"},
			{"label": "PI Edited", "key": "pi_edited"},
			{"label": "Other / Blank status", "key": "other"},
		]},
	]
	return stages


def pipeline_status_map():
	"""{status: row key} for non-round statuses and round select/reject, and
	{status: round} for "in this round" statuses (resolved per candidate by
	pipeline_row_key). Covers every Application Status option plus legacy
	spellings; anything else lands in "other" so the rows add up to Total."""
	meta = frappe.get_meta("Field Registration Form").get_field("application_status")
	statuses = [s.strip() for s in (meta.options or "").split("\n") if s.strip()] if meta else []
	statuses += list(_STATUS_ROW) + ["Round One", "Round Two", "Round Three", "Recruiter Reject",
		"Round 1 Reject", "Round 2 Reject", "Round 3 Reject"]
	to_key, in_round = {}, {}
	for st in statuses:
		if st in _STATUS_ROW:
			to_key[st] = _STATUS_ROW[st]
		elif st in _OFFER_ROWS:
			to_key[st] = "o_" + _slug(st)
		else:
			rnd = canon_round(st)
			if not rnd:
				continue
			res = _round_result(st)
			if res == "in":
				in_round[st] = rnd
			else:
				to_key[st] = _slug(rnd) + "_" + res
	return to_key, in_round


def interview_dates_by_candidate():
	"""{application_id: {round: latest interview date 'YYYY-MM-DD'}} from
	non-cancelled Field Interview Schedules."""
	out = {}
	for r in frappe.db.sql(
		"""select application_id, interview_round, max(interview_date) d
		from `tabField Interview Schedule`
		where ifnull(is_cancelled, 0) = 0 and ifnull(application_id, '') != '' and interview_date is not null
		group by application_id, interview_round""",
		as_dict=True,
	):
		rnd = canon_round(r.interview_round)
		if not rnd:
			continue
		d = str(r.d)
		cur = out.setdefault(r.application_id, {})
		if d > cur.get(rnd, ""):
			cur[rnd] = d
	return out


def pipeline_row_key(status, app_id, to_key, in_round, fis, today_str):
	status = (status or "").strip()
	if status in to_key:
		return to_key[status]
	rnd = in_round.get(status)
	if not rnd:
		return "other"
	dates = fis.get(app_id) or {}
	d = dates.get(rnd)
	if not d:
		# EC / Subject / Technical / Functional are alternatives — an
		# interview booked under a sibling name still counts.
		d = max((v for k, v in dates.items() if ROUND_RANK.get(k) == ROUND_RANK[rnd]), default=None)
	k = _slug(rnd)
	if not d:
		return k + "_pending"
	return k + ("_scheduled" if d >= today_str else "_feedback")


@frappe.whitelist()
def get_report_meta(with_pipeline=0):
	"""Lookups the report pages share: Job Opening locations + the state
	lists for jobStateOf(), and (with_pipeline) the workflow pipeline rows,
	status map and per-candidate interview dates."""
	frappe.has_permission("Field Registration Form", "read", throw=True)
	out = {
		"jo": _job_opening_locations(),
		"states": INDIA_STATES,
		"aliases": STATE_ALIASES,
		"multi": STATE_MULTI,
		"none": STATE_NONE,
		"rounds": [r[0] for r in WORKFLOW_ROUNDS],
	}
	if frappe.utils.cint(with_pipeline):
		to_key, in_round = pipeline_status_map()
		out.update({
			"stages": pipeline_stages(),
			"status_to_key": to_key,
			"in_round": in_round,
			"round_rank": ROUND_RANK,
			"fis": interview_dates_by_candidate(),
			"today": today(),
		})
	return out


# ── Interviewers Data ─────────────────────────────────────────────────────
@frappe.whitelist()
def get_interviewer_data(from_date, to_date):
	"""Rows for the "Interviewers Data" report in reports.js — one row per
	(interview, interviewer) for non-cancelled Field Interview Schedules
	dated from_date..to_date, joined to the candidate's Field Registration
	Form.

	Done in SQL rather than frappe.client.get_list because interviewers
	live in two Table MultiSelect fields (interviewer_email and
	hybrid_interviewers_email) on the same child doctype — get_list's
	child-field syntax joins that child table without telling the two
	fields apart (checked 2026-09-28: both came back identical). Here the
	join is on parentfield explicitly. An interview with no interviewer
	still returns one row, with interviewer NULL.

	Each row also gets:
	  round    — workflow round of the interview (canon_round), "Other" if none
	  selected — 1 if the candidate got past THIS round. A panel member is
	             only credited for the round they sat on (a Leader-round
	             panelist no longer shows Functional-round selections).
	             "Got past" = Select status for the round, a status from a
	             later round / offer stage, or an interview in a later round.
	  _state   — state of the job applied for (job_state)
	"""
	frappe.has_permission("Field Interview Schedule", "read", throw=True)
	rows = frappe.db.sql(
		"""
		select
			s.name, s.application_id, s.applicants_name, s.interview_round, s.interview_date,
			s.role as fis_role, s.department as fis_department, s.location as fis_location,
			c.interviewer_email as interviewer, l.interviewer_name,
			f.name as frf, f.application_status, f.role, f.department, f.location, f.job_title,
			f.date_of_functional_round, f.date_of_final_round
		from `tabField Interview Schedule` s
		left join `tabInterviewer Email Child` c
			on c.parent = s.name and c.parenttype = 'Field Interview Schedule'
			and c.parentfield in ('interviewer_email', 'hybrid_interviewers_email')
		left join `tabInterviewer Email List` l on l.name = c.interviewer_email
		left join `tabField Registration Form` f on f.name = s.application_id
		where ifnull(s.is_cancelled, 0) = 0
			and s.interview_date between %(from_date)s and %(to_date)s
		order by s.interview_date
		""",
		{"from_date": from_date, "to_date": to_date},
		as_dict=True,
	)
	jo = _job_opening_locations()

	# Furthest round each candidate is known to have passed: their status,
	# the round dates on the form, and any later-round interview.
	fis = interview_dates_by_candidate()
	passed = {}
	for r in rows:
		app = r.application_id or r.name
		if app in passed:
			continue
		p = passed_rank(r.application_status)
		if r.date_of_final_round:
			p = max(p, 3)
		elif r.date_of_functional_round:
			p = max(p, 1)
		for rnd in (fis.get(r.application_id) or {}):
			p = max(p, ROUND_RANK[rnd] - 1)
		passed[app] = p

	for r in rows:
		r["round"] = canon_round(r.interview_round) or "Other"
		rank = ROUND_RANK.get(r["round"])
		r["selected"] = 1 if rank and passed[r.application_id or r.name] >= rank else 0
		r["_state"] = (job_state(r.location, jo.get(r.job_title)) if r.frf
			else job_state(r.fis_location))
	return {"rows": rows, "rounds": [w[0] for w in WORKFLOW_ROUNDS]}


# ── Cycle Time ────────────────────────────────────────────────────────────
# Stage dates per candidate. Field Registration Form only stores the
# Applied / Recruiter / Functional / Final round dates (often blank), and
# no Written test or Offer date at all, so each stage takes the first of:
#   1. the date field on the form (when there is one)
#   2. the interview date on Field Interview Schedule for that round
#   3. the day the Application Status first moved into that stage's
#      result statuses (Version history + Field Application Status Log)
_CT_STATUS_STAGE = {"Test Select": "test", "Test Reject": "test", "Test Hold": "test"}
for _n, _sel, _rej, _rk in WORKFLOW_ROUNDS:
	_stage = "recruiter" if _rk == 1 else ("ec" if _rk == 2 else ("final" if _rk >= 4 else None))
	if _stage:
		_CT_STATUS_STAGE[_sel] = _CT_STATUS_STAGE[_rej] = _stage
for _st in _OFFER_ROWS[7:]:  # Offer and everything after it
	_CT_STATUS_STAGE[_st] = "offer"
_VERSION_STATUS_RE = re.compile(r'\["application_status",\s*"[^"]*",\s*"([^"]*)"\]')


def _status_change_dates(since, until):
	"""{frf name: {stage: date}} — earliest status change into each stage
	(latest for "final", the last leader round) between since and until."""
	out = {}

	def add(name, status, d):
		stage = _CT_STATUS_STAGE.get((status or "").strip())
		if not stage or not d:
			return
		d = str(d)[:10]
		cur = out.setdefault(name, {})
		old = cur.get(stage)
		if not old or (d > old if stage == "final" else d < old):
			cur[stage] = d

	for v in frappe.db.sql(
		"""select docname, data, creation from `tabVersion`
		where ref_doctype = 'Field Registration Form' and creation between %s and %s
			and data like %s""",
		(since, until + " 23:59:59", '%"application_status"%'),
		as_dict=True,
	):
		for status in _VERSION_STATUS_RE.findall(v.data or ""):
			add(v.docname, status, v.creation)
	if frappe.db.table_exists("Field Application Status Log"):
		for lg in frappe.db.sql(
			"""select parent, application_status, changed_on from `tabField Application Status Log`
			where parenttype = 'Field Registration Form' and changed_on between %s and %s""",
			(since, until + " 23:59:59"),
			as_dict=True,
		):
			add(lg.parent, lg.application_status, lg.changed_on)
	return out


@frappe.whitelist()
def get_cycle_time_data(from_date, to_date):
	"""Candidates with at least one stage date in from_date..to_date (a
	financial year), each with `d` = {applied, test, recruiter, ec, final,
	offer} as 'YYYY-MM-DD' (missing stages left out) and `_state`."""
	frappe.has_permission("Field Registration Form", "read", throw=True)
	fy_from, fy_to = str(getdate(from_date)), str(getdate(to_date))
	# Look back a year so a stage that started before the FY (e.g. the
	# application) is still found for a transition that ended inside it.
	since = str(getdate(from_date) - timedelta(days=366))

	changes = _status_change_dates(since, fy_to)
	fis = {}
	for r in frappe.db.sql(
		"""select application_id, interview_round, min(interview_date) first_d, max(interview_date) last_d
		from `tabField Interview Schedule`
		where ifnull(is_cancelled, 0) = 0 and ifnull(application_id, '') != ''
			and interview_date between %s and least(%s, curdate())
		group by application_id, interview_round""",
		(since, fy_to),
		as_dict=True,
	):
		rank = ROUND_RANK.get(canon_round(r.interview_round))
		stage = "recruiter" if rank == 1 else ("ec" if rank == 2 else ("final" if rank and rank >= 4 else None))
		if not stage:
			continue
		cur = fis.setdefault(r.application_id, {})
		d = str(r.last_d if stage == "final" else r.first_d)
		old = cur.get(stage)
		if not old or (d > old if stage == "final" else d < old):
			cur[stage] = d

	names = set(changes) | set(fis)
	names.update(frappe.db.sql_list(
		"""select name from `tabField Registration Form`
		where date_of_recruiter_round between %(f)s and %(t)s
			or date_of_functional_round between %(f)s and %(t)s
			or date_of_final_round between %(f)s and %(t)s""",
		{"f": fy_from, "t": fy_to},
	))

	jo = _job_opening_locations()
	out = []
	names = list(names)
	for i in range(0, len(names), 5000):
		for r in frappe.get_all(
			"Field Registration Form",
			filters={"name": ["in", names[i:i + 5000]]},
			fields=["name", "full_name_aadhaar", "application_status", "role", "department", "location",
				"job_title", "date_of_applied", "application_date", "creation", "date_of_recruiter_round",
				"date_of_functional_round", "date_of_final_round"],
		):
			ch, iv = changes.get(r.name, {}), fis.get(r.name, {})
			d = {
				"applied": r.date_of_applied or r.application_date or r.creation,
				"test": ch.get("test"),
				"recruiter": r.date_of_recruiter_round or iv.get("recruiter") or ch.get("recruiter"),
				"ec": r.date_of_functional_round or iv.get("ec") or ch.get("ec"),
				"final": r.date_of_final_round or iv.get("final") or ch.get("final"),
				"offer": ch.get("offer"),
			}
			d = {k: str(v)[:10] for k, v in d.items() if v}
			if not any(fy_from <= v <= fy_to for k, v in d.items() if k != "applied"):
				continue
			out.append({
				"name": r.name, "full_name_aadhaar": r.full_name_aadhaar,
				"application_status": r.application_status, "role": r.role,
				"department": r.department, "location": r.location, "job_title": r.job_title,
				"_state": job_state(r.location, jo.get(r.job_title)), "d": d,
			})
	return out
